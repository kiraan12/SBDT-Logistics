
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, Border, Side
from typing import List, Dict
import io
import csv

class ExcelService:
    @staticmethod
    def generate_shipment_excel(shipments: List[Dict]) -> io.BytesIO:
        """
        Generates Excel file (.xlsx format) in memory matching the strict SBDT format.
        Returns BytesIO buffer containing the .xlsx file.
        """
        wb = Workbook()
        ws = wb.active
        ws.title = "Shipments"

        # Headers matching format.xlsx exactly
        headers = [
            "SL. NO.", "Industry Branch Name", "Tracking Number", "Consignor", "Source",
            "Consignee", "Destination", "Inv No", "Boxes", "Weight",
            "Shipment Type", "Boooking Date ", "Expected Delivery Date", "ETA (From the LSP)",
            "Actual Delivery Date", "Delivery Status", "REMARKS"
        ]

        # Style for header
        header_font = Font(bold=True)
        thin_border = Border(left=Side(style='thin'), right=Side(style='thin'), 
                             top=Side(style='thin'), bottom=Side(style='thin'))

        for col_num, header in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col_num, value=header)
            cell.font = header_font
            cell.border = thin_border
            ws.column_dimensions[cell.column_letter].width = 15

        # Data rows
        for idx, s in enumerate(shipments, 1):
            row_data = [
                idx,                        # SL. NO.
                s.get("branch_name", ""),   # Industry Branch Name
                s.get("lr_no", ""),         # Tracking Number
                s.get("consignor_name", ""), # Consignor
                s.get("source", ""),        # Source
                s.get("consignee_name", ""), # Consignee
                s.get("destination", ""),   # Destination
                s.get("inv_no", ""),        # Inv No
                s.get("boxes", ""),         # Boxes (Invoice Value removed to match format.xlsx)
                s.get("weight", ""),        # Weight
                s.get("shipment_type", ""), # Shipment Type
                str(s.get("booking_date", "")),  # Boooking Date
                str(s.get("expected_delivery_date", "")), # Expected Delivery Date
                str(s.get("eta", "")),      # ETA (From the LSP)
                str(s.get("actual_delivery_date", "")), # Actual Delivery Date
                s.get("delivery_status", ""), # Delivery Status
                s.get("remarks", "")       # REMARKS
            ]

            for col_num, value in enumerate(row_data, 1):
                cell = ws.cell(row=idx + 1, column=col_num, value=value)
                cell.border = thin_border

        output = io.BytesIO()
        # Save as .xlsx format (openpyxl always saves as .xlsx)
        wb.save(output)
        output.seek(0)  # Reset buffer position for reading
        return output

    @staticmethod
    def generate_shipment_csv(shipments: List[Dict]) -> io.BytesIO:
        """
        Generates CSV file in memory matching the SBDT format.
        Returns BytesIO buffer containing the CSV file.
        """
        output = io.BytesIO()
        
        # Headers matching format.xlsx exactly
        headers = [
            "SL. NO.", "Industry Branch Name", "Tracking Number", "Consignor", "Source",
            "Consignee", "Destination", "Inv No", "Boxes", "Weight",
            "Shipment Type", "Boooking Date ", "Expected Delivery Date", "ETA (From the LSP)",
            "Actual Delivery Date", "Delivery Status", "REMARKS"
        ]
        
        # Use StringIO for text, then encode to bytes
        text_buffer = io.StringIO()
        writer = csv.writer(text_buffer)
        
        # Write headers
        writer.writerow(headers)
        
        # Write data rows
        for idx, s in enumerate(shipments, 1):
            row_data = [
                idx,                        # SL. NO.
                s.get("branch_name", ""),   # Industry Branch Name
                s.get("lr_no", ""),         # Tracking Number
                s.get("consignor_name", ""), # Consignor
                s.get("source", ""),        # Source
                s.get("consignee_name", ""), # Consignee
                s.get("destination", ""),   # Destination
                s.get("inv_no", ""),        # Inv No
                s.get("boxes", ""),         # Boxes
                s.get("weight", ""),        # Weight
                s.get("shipment_type", ""), # Shipment Type
                str(s.get("booking_date", "")),  # Boooking Date
                str(s.get("expected_delivery_date", "")), # Expected Delivery Date
                str(s.get("eta", "")),      # ETA (From the LSP)
                str(s.get("actual_delivery_date", "")), # Actual Delivery Date
                s.get("delivery_status", ""), # Delivery Status
                s.get("remarks", "")       # REMARKS
            ]
            writer.writerow(row_data)
        
        # Convert string buffer to bytes
        csv_content = text_buffer.getvalue()
        output.write(csv_content.encode('utf-8-sig'))  # UTF-8 with BOM for Excel compatibility
        output.seek(0)  # Reset buffer position for reading
        return output
